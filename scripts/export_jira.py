"""Export one Jira project into plain-text files the sandbox can read.

Runs on the local machine, which can reach Jira; the sandbox cannot. Output is
one folder per ticket under data/<project>/:

    data/vxep/VXEP-123/ticket.md                 every non-empty field, all comments,
                                                 and a table of every attachment
    data/vxep/VXEP-123/attachments/<name>.md     text pulled from one document

plus data/<project>/index.jsonl (one line per ticket) and manifest.json (what
ran, and every attachment that could not be converted, with the reason).

Documents (pdf, docx, pptx, xlsx, txt, msg, eml, ...) are converted to text.
CSVs get their column headers and row count only. AAR packages, videos and
images are listed in ticket.md but not downloaded (decisions, October 20).

Downloaded originals are kept in --cache, outside the repo, so a re-run only
downloads what is new.

    .venv/bin/python scripts/export_jira.py --env ../jira-dash/.env \
        --cache ../root_cause_local/attachments --project VXEP

The env file needs JIRA_URL, JIRA_EMAIL and JIRA_API_TOKEN. Nothing falls back
to a default: a missing setting stops the run.
"""
import argparse
import csv
import email
import email.policy
import hashlib
import io
import json
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone

import requests
from dotenv import dotenv_values

# Converted text longer than this is cut, with a note saying so, so one huge
# manual cannot swamp a ticket. Plain-text attachments are mostly XCOMM/VTOL
# event logs, where the failure is usually at the end, and they total only
# ~17 MB across VXEP, so they get a far higher limit.
MAX_CHARS = 150_000
MAX_CHARS_TEXT = 5_000_000

# Custom fields that are Jira bookkeeping, not ticket content.
SKIP_FIELD_NAMES = {'Rank', 'Development'}
SKIP_FIELD_PREFIXES = ('[CHART]',)
# A long value repeated on this share of tickets is template text (sign-off
# wording, instructions), not something written about this ticket.
BOILERPLATE_SHARE = 0.2
BOILERPLATE_MIN_CHARS = 60
# Spreadsheet rows written per sheet.
MAX_SHEET_ROWS = 300

DOC_KINDS = {
    'pdf': 'pdf', 'docx': 'docx', 'pptx': 'pptx', 'xlsx': 'xlsx', 'xlsm': 'xlsx',
    'msg': 'msg', 'eml': 'eml',
    'txt': 'text', 'md': 'text', 'json': 'text', 'jsonl': 'text', 'log': 'text',
    'html': 'html', 'htm': 'html', 'xml': 'text', 'py': 'text', 'sql': 'text',
}
CSV_KINDS = {'csv'}
IMAGE_KINDS = {'png', 'jpg', 'jpeg', 'jfif', 'heic', 'tif', 'tiff', 'gif', 'bmp', 'webp'}
VIDEO_KINDS = {'mp4', 'mov', 'mpg', 'mpeg', 'avi', 'ts', 'mkv', 'wmv', 'm4v'}
ARCHIVE_KINDS = {'zip', '7z', 'gz', 'tgz', 'tar', 'rar', 'changetozip'}

# System fields worth writing out. Everything else from Jira's system set is
# bookkeeping (votes, watches, rank, avatars...). All non-empty custom fields
# are written, under their display names.
SYSTEM_FIELDS = [
    ('summary', 'Summary'), ('issuetype', 'Issue type'), ('status', 'Status'),
    ('resolution', 'Resolution'), ('priority', 'Priority'), ('created', 'Created'),
    ('updated', 'Updated'), ('resolutiondate', 'Resolved'), ('duedate', 'Due'),
    ('reporter', 'Reporter'), ('assignee', 'Assignee'), ('creator', 'Creator'),
    ('labels', 'Labels'), ('components', 'Components'), ('versions', 'Affects versions'),
    ('fixVersions', 'Fix versions'), ('environment', 'Environment'), ('parent', 'Parent'),
    ('issuelinks', 'Links'),
]
LONG_SYSTEM_FIELDS = [('description', 'Description')]


# ---------------------------------------------------------------------------
# Jira
# ---------------------------------------------------------------------------

class Jira:
    def __init__(self, env_path):
        env = dotenv_values(env_path)
        missing = [k for k in ('JIRA_URL', 'JIRA_EMAIL', 'JIRA_API_TOKEN') if not env.get(k)]
        if missing:
            sys.exit('Missing in %s: %s' % (env_path, ', '.join(missing)))
        self.url = env['JIRA_URL'].rstrip('/')
        self.s = requests.Session()
        self.s.auth = (env['JIRA_EMAIL'], env['JIRA_API_TOKEN'])
        self.s.headers['Accept'] = 'application/json'

    def get(self, path, **params):
        for attempt in range(5):
            r = self.s.get(self.url + path, params=params, timeout=60)
            if r.status_code == 429 or r.status_code >= 500:
                time.sleep(int(r.headers.get('Retry-After', 2 ** attempt)))
                continue
            r.raise_for_status()
            return r.json()
        r.raise_for_status()

    def search(self, jql):
        issues, token = [], None
        while True:
            body = {'jql': jql, 'maxResults': 100, 'fields': ['*all']}
            if token:
                body['nextPageToken'] = token
            r = self.s.post(self.url + '/rest/api/3/search/jql', json=body, timeout=120)
            r.raise_for_status()
            data = r.json()
            issues += data.get('issues', [])
            token = data.get('nextPageToken')
            print('  %d tickets' % len(issues), file=sys.stderr)
            if not token:
                return issues

    def all_comments(self, key, field):
        """The search returns only some comments on busy tickets; fetch the rest."""
        comments = list((field or {}).get('comments') or [])
        total = (field or {}).get('total', len(comments))
        start = len(comments)
        while start < total:
            page = self.get('/rest/api/3/issue/%s/comment' % key, startAt=start, maxResults=100)
            got = page.get('comments') or []
            if not got:
                break
            comments += got
            start += len(got)
        return comments

    def download(self, url):
        for attempt in range(5):
            r = self.s.get(url, timeout=600, headers={'Accept': '*/*'})
            if r.status_code == 429 or r.status_code >= 500:
                time.sleep(int(r.headers.get('Retry-After', 2 ** attempt)))
                continue
            r.raise_for_status()
            return r.content
        r.raise_for_status()


# ---------------------------------------------------------------------------
# Rendering Jira values as text
# ---------------------------------------------------------------------------

def adf_to_text(node):
    """Atlassian Document Format -> plain text, keeping lists, tables and breaks."""
    if node is None:
        return ''
    if isinstance(node, str):
        return node
    if isinstance(node, list):
        return ''.join(adf_to_text(n) for n in node)
    if not isinstance(node, dict):
        return ''
    t = node.get('type')
    attrs = node.get('attrs') or {}
    if t == 'text':
        return node.get('text', '')
    if t == 'mention':
        return attrs.get('text', '')
    if t in ('emoji',):
        return attrs.get('text') or attrs.get('shortName', '')
    if t == 'hardBreak':
        return '\n'
    if t in ('inlineCard', 'blockCard', 'embedCard'):
        return attrs.get('url', '')
    if t in ('media', 'mediaInline'):
        return '[embedded %s]' % (attrs.get('alt') or attrs.get('type') or 'media')
    if t == 'date':
        try:
            return datetime.fromtimestamp(int(attrs.get('timestamp')) / 1000, timezone.utc).strftime('%Y-%m-%d')
        except (TypeError, ValueError):
            return ''
    if t == 'status':
        return '[%s]' % attrs.get('text', '')
    inner = adf_to_text(node.get('content', []))
    if t == 'listItem':
        return '- ' + inner.strip() + '\n'
    if t == 'tableCell' or t == 'tableHeader':
        return inner.strip().replace('\n', ' ') + ' | '
    if t == 'tableRow':
        return '| ' + inner.rstrip() + '\n'
    if t == 'codeBlock':
        return '```\n' + inner + '\n```\n'
    if t == 'heading':
        return '#' * int(attrs.get('level', 3)) + ' ' + inner + '\n\n'
    if t in ('paragraph', 'blockquote', 'panel', 'bulletList', 'orderedList', 'table', 'mediaSingle', 'mediaGroup', 'expand', 'nestedExpand'):
        return inner + '\n'
    return inner


def value_to_text(v):
    """A Jira field value of any shape -> one string. People fields become
    display names only (addresses typed into text are left as written)."""
    if v is None or v == '' or v == [] or v == {}:
        return ''
    if isinstance(v, str):
        return v.strip()
    if isinstance(v, (int, float, bool)):
        return str(v)
    if isinstance(v, list):
        return ', '.join(x for x in (value_to_text(i) for i in v) if x)
    if isinstance(v, dict):
        if v.get('type') == 'doc':
            return adf_to_text(v).strip()
        if 'accountId' in v:
            return v.get('displayName', '')
        if 'value' in v:
            child = value_to_text(v.get('child'))
            return v['value'] + (' / ' + child if child else '')
        if 'key' in v and 'fields' in v:      # parent
            return '%s %s' % (v['key'], (v['fields'] or {}).get('summary', ''))
        for k in ('name', 'displayName', 'title'):
            if v.get(k):
                return str(v[k])
    return ''


def links_to_text(links):
    out = []
    for link in links or []:
        for direction, word in (('outwardIssue', 'outward'), ('inwardIssue', 'inward')):
            other = link.get(direction)
            if other:
                verb = (link.get('type') or {}).get(word, 'relates to')
                out.append('%s %s (%s)' % (verb, other.get('key'), (other.get('fields') or {}).get('summary', '')))
    return '; '.join(out)


# ---------------------------------------------------------------------------
# Attachments
# ---------------------------------------------------------------------------

def kind_of(filename):
    ext = filename.lower().rsplit('.', 1)[-1] if '.' in filename else ''
    if ext in DOC_KINDS:
        return DOC_KINDS[ext]
    if ext in CSV_KINDS:
        return 'csv'
    if ext in IMAGE_KINDS:
        return 'image'
    if ext in VIDEO_KINDS:
        return 'video'
    if ext in ARCHIVE_KINDS:
        return 'archive'
    return 'other'


def decode_text(data):
    for enc in ('utf-8-sig', 'utf-16') if data[:2] in (b'\xff\xfe', b'\xfe\xff') else ('utf-8-sig',):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode('cp1252', errors='replace')


def convert_pdf(data):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or '').strip()
        pages.append('## Page %d\n\n%s' % (i, text or '(no text on this page; probably an image or scan)'))
    return '\n\n'.join(pages), {'pages': len(reader.pages)}


def _docx_table(table):
    rows = []
    for row in table.rows:
        rows.append('| ' + ' | '.join(c.text.strip().replace('\n', ' ') for c in row.cells) + ' |')
    return '\n'.join(rows)


def convert_docx(data):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(io.BytesIO(data))
    parts = []
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            p = Paragraph(block, d)
            text = p.text.strip()
            if text:
                style = (p.style.name or '').lower() if p.style is not None else ''
                parts.append(('## ' if style.startswith('heading') else '') + text)
        elif tag == 'tbl':
            parts.append(_docx_table(Table(block, d)))
    images = len([r for r in d.part.rels.values() if 'image' in r.reltype])
    return '\n\n'.join(parts), {'images': images}


def convert_pptx(data):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(io.BytesIO(data))
    slides, images = [], 0

    def shape_text(shape):
        nonlocal images
        out = []
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for s in shape.shapes:
                out += shape_text(s)
            return out
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            images += 1
            out.append('[picture]')
        if getattr(shape, 'has_text_frame', False) and shape.text_frame.text.strip():
            out.append(shape.text_frame.text.strip())
        if getattr(shape, 'has_table', False) and shape.has_table:
            for row in shape.table.rows:
                out.append('| ' + ' | '.join(c.text.strip().replace('\n', ' ') for c in row.cells) + ' |')
        if getattr(shape, 'has_chart', False) and shape.has_chart:
            chart = shape.chart
            title = chart.chart_title.text_frame.text if chart.has_title else ''
            out.append('[chart%s]' % (': ' + title if title else ''))
        return out

    for i, slide in enumerate(prs.slides, 1):
        lines = []
        for shape in slide.shapes:
            lines += shape_text(shape)
        notes = ''
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        text = '## Slide %d\n\n%s' % (i, '\n'.join(lines) or '(no text)')
        if notes:
            text += '\n\nSpeaker notes:\n' + notes
        slides.append(text)
    return '\n\n'.join(slides), {'slides': len(prs.slides), 'images': images}


def convert_xlsx(data):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows, total = [], 0
        for row in ws.iter_rows(values_only=True):
            if not any(c is not None and str(c).strip() for c in row):
                continue
            total += 1
            if len(rows) < MAX_SHEET_ROWS:
                rows.append('\t'.join('' if c is None else str(c).replace('\n', ' ') for c in row).rstrip())
        text = '## Sheet: %s\n\n%s' % (ws.title, '\n'.join(rows) or '(empty)')
        if total > MAX_SHEET_ROWS:
            text += '\n\n(cut: %d of %d non-empty rows shown)' % (MAX_SHEET_ROWS, total)
        parts.append(text)
    return '\n\n'.join(parts), {'sheets': len(wb.worksheets)}


def _email_text(sender, to, date, subject, body):
    head = '\n'.join('%s: %s' % (k, v) for k, v in
                     (('From', sender), ('To', to), ('Date', date), ('Subject', subject)) if v)
    return head + '\n\n' + (body or '').strip()


def _strip_addresses(s):
    """Keep names, drop email addresses: 'Jo Smith <jo@x.com>' -> 'Jo Smith'."""
    s = re.sub(r'\s*<[^>]*@[^>]*>', '', s or '')
    return re.sub(r'[\w.+-]+@[\w-]+\.[\w.-]+', '[email]', s)


def convert_msg(data):
    import extract_msg
    m = extract_msg.openMsg(io.BytesIO(data))
    try:
        body = m.body or ''
        text = _email_text(_strip_addresses(m.sender), _strip_addresses(m.to), str(m.date or ''),
                           m.subject, body)
        names = [a.longFilename or a.shortFilename or '' for a in m.attachments]
    finally:
        m.close()
    if names:
        text += '\n\nAttached to this email (not converted): ' + ', '.join(n for n in names if n)
    return text, {}


def convert_eml(data):
    msg = email.message_from_bytes(data, policy=email.policy.default)
    part = msg.get_body(preferencelist=('plain', 'html'))
    body = part.get_content() if part else ''
    if part is not None and part.get_content_type() == 'text/html':
        body = html_to_text(body)
    names = [p.get_filename() for p in msg.iter_attachments() if p.get_filename()]
    text = _email_text(_strip_addresses(str(msg['from'] or '')), _strip_addresses(str(msg['to'] or '')),
                       str(msg['date'] or ''), str(msg['subject'] or ''), body)
    if names:
        text += '\n\nAttached to this email (not converted): ' + ', '.join(names)
    return text, {}


def html_to_text(html):
    from bs4 import BeautifulSoup
    return BeautifulSoup(html, 'html.parser').get_text('\n')


def convert_html(data):
    return html_to_text(decode_text(data)), {}


def convert_text(data):
    return decode_text(data), {}


def csv_summary(data):
    text = decode_text(data)
    reader = csv.reader(io.StringIO(text))
    header = next(reader, [])
    rows = sum(1 for _ in reader)
    return ('Columns (%d):\n%s\n\nData rows: %d (data not exported)'
            % (len(header), '\n'.join('- ' + h for h in header), rows)), {'rows': rows, 'columns': len(header)}


CONVERTERS = {
    'pdf': convert_pdf, 'docx': convert_docx, 'pptx': convert_pptx, 'xlsx': convert_xlsx,
    'msg': convert_msg, 'eml': convert_eml, 'html': convert_html, 'text': convert_text,
    'csv': csv_summary,
}


def safe_name(name):
    """Windows-safe file name, unique enough to sit next to its siblings."""
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name).strip(' .')
    return name[:150] or 'attachment'


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def fetch_attachment(jira, att, cache):
    path = os.path.join(cache, '%s_%s' % (att['id'], safe_name(att['filename'])))
    if os.path.exists(path) and os.path.getsize(path) == att.get('size', -1):
        with open(path, 'rb') as f:
            return f.read()
    data = jira.download(att['content'])
    tmp = path + '.part'
    with open(tmp, 'wb') as f:
        f.write(data)
    os.replace(tmp, path)
    return data


def process_attachment(jira, att, cache, out_dir):
    """Returns a dict describing what happened. Never raises."""
    kind = kind_of(att['filename'])
    info = {'id': att['id'], 'filename': att['filename'], 'kind': kind,
            'size_kb': round(att.get('size', 0) / 1024), 'created': (att.get('created') or '')[:10],
            'author': value_to_text(att.get('author'))}
    if kind not in CONVERTERS:
        info['result'] = 'listed only'
        return info
    try:
        data = fetch_attachment(jira, att, cache)
        text, extra = CONVERTERS[kind](data)
        info.update(extra)
        text = text.strip()
        if not text:
            info['result'] = 'no text found'
            return info
        limit = MAX_CHARS_TEXT if kind == 'text' else MAX_CHARS
        cut = len(text) > limit
        if cut:
            text = text[:limit] + '\n\n(cut at %d characters of %d)' % (limit, len(text))
        name = '%s_%s.md' % (att['id'], safe_name(att['filename']))
        with open(os.path.join(out_dir, name), 'w', encoding='utf-8', newline='\n') as f:
            f.write('# %s\n\nAttachment %s, %s, %s KB, added %s by %s.\n\n%s\n'
                    % (att['filename'], att['id'], kind, info['size_kb'], info['created'],
                       info['author'] or 'unknown', text))
        info['result'] = 'converted (cut)' if cut else 'converted'
        info['text_file'] = 'attachments/' + name
    except Exception as e:                       # recorded, never silent
        info['result'] = 'failed'
        info['error'] = '%s: %s' % (type(e).__name__, str(e)[:300])
    return info


def boilerplate_values(issues, field_names):
    """{field id: set of values} that are template text, by how often they repeat."""
    seen = {}
    for issue in issues:
        for fid, v in issue['fields'].items():
            if fid.startswith('customfield_') or fid == 'description':
                text = value_to_text(v)
                if len(text) >= BOILERPLATE_MIN_CHARS:
                    seen.setdefault(fid, {}).setdefault(text, 0)
                    seen[fid][text] += 1
    floor = max(5, BOILERPLATE_SHARE * len(issues))
    return {fid: {t for t, n in counts.items() if n >= floor} for fid, counts in seen.items()}


def skip_field(name):
    return name in SKIP_FIELD_NAMES or name.startswith(SKIP_FIELD_PREFIXES)


def ticket_markdown(issue, field_names, comments, attachments, boilerplate):
    f = issue['fields']
    key = issue['key']
    lines = ['# %s: %s' % (key, (f.get('summary') or '').strip()), '']
    for fid, label in SYSTEM_FIELDS:
        if fid == 'summary':
            continue
        text = links_to_text(f.get(fid)) if fid == 'issuelinks' else value_to_text(f.get(fid))
        if text:
            if fid in ('created', 'updated', 'resolutiondate'):
                text = text[:10]
            lines.append('- **%s:** %s' % (label, text))
    long_fields = []
    for fid in sorted(k for k in f if k.startswith('customfield_')):
        text = value_to_text(f.get(fid))
        label = field_names.get(fid, fid)
        if not text or skip_field(label) or text in boilerplate.get(fid, ()):
            continue
        if '\n' in text or len(text) > 120:
            long_fields.append((label, text))
        else:
            lines.append('- **%s:** %s' % (label, text))
    for fid, label in LONG_SYSTEM_FIELDS:
        text = value_to_text(f.get(fid))
        if text and text not in boilerplate.get(fid, ()):
            long_fields.insert(0, (label, text))
    for label, text in long_fields:
        lines += ['', '## %s' % label, '', text]

    lines += ['', '## Comments (%d)' % len(comments), '']
    for c in comments:
        body = adf_to_text(c.get('body')).strip()
        lines += ['### %s, %s' % (value_to_text(c.get('author')) or 'unknown', (c.get('created') or '')[:10]),
                  '', body or '(empty)', '']

    lines += ['## Attachments (%d)' % len(attachments), '']
    if attachments:
        lines += ['| File | Kind | KB | Added | Result |', '| --- | --- | --- | --- | --- |']
        for a in attachments:
            name = a['filename'].replace('|', '/')
            cell = '[%s](%s)' % (name, a['text_file'].replace(' ', '%20')) if a.get('text_file') else name
            result = a['result'] + (': ' + a['error'] if a.get('error') else '')
            lines.append('| %s | %s | %s | %s | %s |' % (cell, a['kind'], a['size_kb'], a['created'], result.replace('|', '/')))
    return '\n'.join(lines).rstrip() + '\n'


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('\n\n')[0])
    ap.add_argument('--env', required=True, help='dotenv file with JIRA_URL, JIRA_EMAIL, JIRA_API_TOKEN')
    ap.add_argument('--cache', required=True, help='folder outside the repo for downloaded originals')
    ap.add_argument('--project', required=True, help='Jira project key, e.g. VXEP')
    ap.add_argument('--out', default=None, help='output folder (default data/<project>)')
    ap.add_argument('--workers', type=int, default=4)
    ap.add_argument('--jql', default=None, help='extra JQL to narrow the run, for testing')
    args = ap.parse_args(argv)

    repo = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    out_root = args.out or os.path.join(repo, 'data', args.project.lower())
    if os.path.abspath(args.cache).startswith(repo + os.sep):
        sys.exit('--cache must be outside the repository: originals are never committed')
    os.makedirs(args.cache, exist_ok=True)
    os.makedirs(out_root, exist_ok=True)

    jira = Jira(args.env)
    field_names = {f['id']: f['name'] for f in jira.get('/rest/api/3/field')}
    print('Searching %s ...' % args.project, file=sys.stderr)
    jql = 'project = %s' % args.project
    if args.jql:
        jql += ' AND (%s)' % args.jql
    issues = jira.search(jql + ' ORDER BY key ASC')

    boilerplate = boilerplate_values(issues, field_names)
    started = datetime.now(timezone.utc)
    index, failures, counts = [], [], {}
    pool = ThreadPoolExecutor(max_workers=args.workers)
    for n, issue in enumerate(issues, 1):
        key, f = issue['key'], issue['fields']
        tdir = os.path.join(out_root, key)
        adir = os.path.join(tdir, 'attachments')
        atts = f.get('attachment') or []
        if any(kind_of(a['filename']) in CONVERTERS for a in atts):
            os.makedirs(adir, exist_ok=True)
        else:
            os.makedirs(tdir, exist_ok=True)
        comments = jira.all_comments(key, f.get('comment'))
        results = list(pool.map(lambda a: process_attachment(jira, a, args.cache, adir), atts))
        results.sort(key=lambda a: (a['created'], a['filename']))
        for a in results:
            counts[a['result']] = counts.get(a['result'], 0) + 1
            if a['result'] == 'failed':
                failures.append({'ticket': key, 'filename': a['filename'], 'error': a['error']})
        with open(os.path.join(tdir, 'ticket.md'), 'w', encoding='utf-8', newline='\n') as fh:
            fh.write(ticket_markdown(issue, field_names, comments, results, boilerplate))
        index.append({
            'key': key, 'summary': f.get('summary'), 'status': value_to_text(f.get('status')),
            'issuetype': value_to_text(f.get('issuetype')),
            'created': (f.get('created') or '')[:10], 'resolved': (f.get('resolutiondate') or '')[:10],
            'labels': f.get('labels') or [], 'components': value_to_text(f.get('components')),
            'comments': len(comments),
            'attachments': {k: sum(1 for a in results if a['kind'] == k)
                            for k in sorted({a['kind'] for a in results})},
            'converted': sum(1 for a in results if a.get('text_file')),
        })
        if n % 50 == 0:
            print('  %d / %d tickets written' % (n, len(issues)), file=sys.stderr)
    pool.shutdown()

    with open(os.path.join(out_root, 'index.jsonl'), 'w', encoding='utf-8', newline='\n') as fh:
        for row in index:
            fh.write(json.dumps(row, ensure_ascii=False) + '\n')
    manifest = {
        'project': args.project, 'exported_at': started.isoformat(timespec='seconds'),
        'tickets': len(issues), 'attachment_results': counts, 'failures': failures,
        'max_chars_per_attachment': MAX_CHARS, 'max_chars_per_text_attachment': MAX_CHARS_TEXT,
        'max_rows_per_sheet': MAX_SHEET_ROWS,
        'boilerplate_skipped': {field_names.get(k, k): len(v) for k, v in boilerplate.items() if v},
    }
    with open(os.path.join(out_root, 'manifest.json'), 'w', encoding='utf-8', newline='\n') as fh:
        json.dump(manifest, fh, indent=2, ensure_ascii=False)
    print(json.dumps({k: manifest[k] for k in ('tickets', 'attachment_results')}, indent=2))
    print('%d attachment failures; see manifest.json' % len(failures))


if __name__ == '__main__':
    main()
